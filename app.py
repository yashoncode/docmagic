"""DocMagic / XEON AI — chat with your Docs (PDFs and Excel). LangChain + Streamlit edition.

RAG pipeline on LangChain abstractions with a Streamlit frontend:
loaders -> splitter -> Chroma vector store -> LCEL chain.
Run: streamlit run app.py
"""

import ast
import asyncio
import io
import json
import logging
import operator
import os
import tempfile
from datetime import datetime
from operator import itemgetter
from urllib.parse import urlparse
from uuid import uuid4

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_community.document_loaders import PyPDFLoader
from langchain_community.retrievers import BM25Retriever
from langchain_core.documents import Document
from langchain_core.output_parsers import JsonOutputParser, StrOutputParser
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_core.runnables import RunnableLambda, RunnableParallel, RunnablePassthrough
from langchain_nvidia_ai_endpoints import NVIDIAEmbeddings, NVIDIARerank
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openai import AuthenticationError
from openpyxl import load_workbook

load_dotenv(override=True)  # .env edits win on rerun

# root stays at WARNING so httpx/streamlit chatter is hidden; only our logs show
logging.basicConfig(format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("docmagic")
log.setLevel(logging.INFO)

LLM_BASE_URL = os.getenv("LLM_BASE_URL", "https://integrate.api.nvidia.com/v1")
LLM_API_KEY = os.getenv("LLM_API_KEY", "")
LLM_MODEL = os.getenv("LLM_MODEL", "nvidia/nemotron-3-super-120b-a12b")
EMBED_API_KEY = os.getenv("EMBED_API_KEY") or LLM_API_KEY
# embeddings keep their own endpoint so switching LLM provider can't break them
EMBED_BASE_URL = os.getenv("EMBED_BASE_URL", "https://integrate.api.nvidia.com/v1")
EMBED_MODEL = os.getenv("EMBED_MODEL", "nvidia/nemotron-3-embed-1b")
RERANK_MODEL = os.getenv("RERANK_MODEL", "nvidia/llama-nemotron-rerank-vl-1b-v2")  # reranks hybrid hits
RERANK_API_KEY = os.getenv("RERANK_API_KEY", "")  # own key; falls back to the embed key

CHUNK_CHARS = 1500  # splitter counts characters; ~250 words
CHUNK_OVERLAP = 300  # ~50 words
TOP_K = 5
FETCH_K = 20  # hybrid candidates pulled per retriever before reranking down to TOP_K
WEB_RESULTS = 4  # web results pulled per question when web search is on
MAX_FILES = 2
SUMMARY_WORDS = 1500  # per file; keeps the summary prompt small and cheap
CHART_SERIES = 3  # max numeric series per auto-chart; more gets unreadable

CURATED_MODELS = [
    "nvidia/nemotron-3-super-120b-a12b",
    "qwen/qwen3-next-80b-a3b-instruct",
    "meta/llama-3.3-70b-instruct",
    "meta/llama-3.1-8b-instruct",
    "deepseek-ai/deepseek-v4-flash",
]
CUSTOM_MODEL = "Custom…"

SUGGESTIONS = {
    ":blue[:material/help:] What is cross-docking?": "What is cross-docking?",
    ":orange[:material/school:] Incoterms in simple terms": "Explain Incoterms in simple terms",
    ":violet[:material/description:] What's in my documents?": (
        "Give me an overview of what my documents contain"
    ),
}

SYSTEM_PROMPT = (
    "You are XEON AI, a friendly expert on logistics, warehousing, ERP and CRM "
    "documents (SOPs, rate cards, contracts, manuals). Talk like a helpful "
    "colleague: warm, plain language, short sentences — never stiff or robotic. "
    "For questions about the documents, answer ONLY from the provided context and "
    "cite the source like [file.pdf p.3] or [file.xlsx Rates] (sheet name for "
    "spreadsheets). The context may include live web results labeled [web ...]; "
    "prefer the user's documents for specifics and use web results for general or "
    "current information, citing them by domain. If the context does not "
    "have the answer, say so in one friendly sentence — do not lecture about what "
    "the context contains. "
    "If the user is just greeting or reacting ('hi', 'nice', 'thanks'), reply "
    "warmly in a sentence or two — no citations, no mention of context or documents."
)

AGENT_PROMPT = SYSTEM_PROMPT + (
    " You have tools. Use search_documents for anything about the user's uploaded files. "
    "Use search_web for general knowledge, logistics concepts or current information. "
    "Use calculate for ALL arithmetic — rates, totals, percentages, GST — never do math "
    "in your head; fetch the numbers first, then calculate. Show the working briefly and "
    "cite document excerpts by file and web results by domain."
)

# splits on paragraphs first, then lines, then words — chunks end at natural
# boundaries instead of mid-sentence
splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_CHARS, chunk_overlap=CHUNK_OVERLAP)

ANSWER_PROMPT = ChatPromptTemplate.from_messages(
    [
        ("system", SYSTEM_PROMPT),
        MessagesPlaceholder("history"),
        ("human", "Context:\n{context}\n\nQuestion: {question}"),
    ]
)

CHART_SPEC_PROMPT = ChatPromptTemplate.from_messages(
    [
        (
            "system",
            "Translate the user's chart request into JSON. Available columns (with dtype): "
            "{columns}. Reply with ONLY this JSON, no prose: "
            '{{"chart": "bar|line|area|scatter|pie", "x": "<column or null>", '
            '"y": ["<numeric columns>"], "agg": "sum|mean|count|none", "top_n": 12}}. '
            "Choose the type that fits the data: change over time -> line or area, "
            "category comparison -> bar, share of a whole with few categories -> pie, "
            "relationship between two numeric columns -> scatter.",
        ),
        ("human", "{question}"),
    ]
)

SUMMARY_PROMPT = ChatPromptTemplate.from_messages(
    [
        ("system", SYSTEM_PROMPT),
        (
            "human",
            "Give a short plain-language summary of each document below "
            "(a few bullet points per document), then one line on what they cover "
            "together.\n\n{excerpts}",
        ),
    ]
)


def secret(name: str) -> str:
    """st.secrets first (Streamlit Cloud), then environment (.env locally)."""
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:  # no secrets.toml anywhere
        pass
    return os.getenv(name, "")


@st.cache_resource(max_entries=4)
def resources(llm_key: str, llm_base_url: str, llm_model: str, embed_key: str):
    """LLM + embeddings, cached per (key, model) so sidebar changes take effect."""
    kwargs = {}
    if "nemotron" in llm_model:
        kwargs["extra_body"] = {"chat_template_kwargs": {"enable_thinking": False}}
    if "kimi" not in llm_model:
        kwargs["temperature"] = 0.2  # kimi-k3 fixes temperature=1.0 and rejects overrides
    llm = ChatOpenAI(base_url=llm_base_url, api_key=llm_key, model=llm_model, **kwargs)
    # embed_documents() sends input_type=passage, embed_query() sends query
    embeddings = NVIDIAEmbeddings(
        base_url=EMBED_BASE_URL, model=EMBED_MODEL, api_key=embed_key, truncate="END"
    )
    return llm, embeddings


# in-process calculator (Cloud-safe: no subprocess, no MCP server to spawn)
_CALC_OPS = {
    ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
    ast.Div: operator.truediv, ast.FloorDiv: operator.floordiv, ast.Mod: operator.mod,
    ast.Pow: operator.pow, ast.USub: operator.neg, ast.UAdd: operator.pos,
}


def safe_calc(expression: str) -> float:
    """Evaluate arithmetic safely — numbers and + - * / // % ** only, never runs code."""

    def _eval(node):
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.BinOp) and type(node.op) in _CALC_OPS:
            return _CALC_OPS[type(node.op)](_eval(node.left), _eval(node.right))
        if isinstance(node, ast.UnaryOp) and type(node.op) in _CALC_OPS:
            return _CALC_OPS[type(node.op)](_eval(node.operand))
        raise ValueError("unsupported expression")

    return _eval(ast.parse(expression, mode="eval").body)


def docs_store(embeddings) -> Chroma:
    """Per-browser-session docs collection: visitors never see each other's files."""
    sid = st.session_state.setdefault("session_id", uuid4().hex[:12])
    # ponytail: orphaned session collections pile up until restart; fine on ephemeral cloud
    return Chroma(
        collection_name=f"docs-{sid}", embedding_function=embeddings, persist_directory="chroma_db"
    )


def default_resources():
    """Env-configured resources, for scripts and tests."""
    return resources(LLM_API_KEY, LLM_BASE_URL, LLM_MODEL, EMBED_API_KEY)


def load_units(path: str) -> list[Document]:
    """One Document per PDF page or Excel sheet, labeled for citations."""
    name = os.path.basename(path)
    if path.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(path, read_only=True, data_only=True)
        docs = []
        for ws in wb.worksheets:
            rows = (
                " | ".join(str(c) for c in row if c is not None)
                for row in ws.iter_rows(values_only=True)
            )
            text = "\n".join(r for r in rows if r)
            docs.append(Document(page_content=text, metadata={"source": name, "loc": ws.title}))
        wb.close()
        return docs
    return [
        Document(
            page_content=p.page_content,
            metadata={"source": name, "loc": f"p.{p.metadata.get('page', 0) + 1}"},
        )
        for p in PyPDFLoader(path).load()
    ]


def ingest(paths: list[str], vectorstore: Chroma) -> int:
    """Load, split and index the documents; returns the chunk count."""
    chunks = splitter.split_documents(doc for path in paths for doc in load_units(path))
    if chunks:
        vectorstore.add_documents(chunks)  # NVIDIAEmbeddings batches 50/call itself
    return len(chunks)


def format_docs(docs: list[Document]) -> str:
    log.info("context <- %s", [f"{d.metadata['source']} {d.metadata['loc']}" for d in docs])
    return "\n\n".join(
        f"[{d.metadata['source']} {d.metadata['loc']}]\n{d.page_content}" for d in docs
    )


@st.cache_resource
def reranker(api_key: str):
    """NVIDIA NIM cross-encoder reranker; runs on the app's free embed key/endpoint."""
    return NVIDIARerank(model=RERANK_MODEL, base_url=EMBED_BASE_URL, api_key=api_key, top_n=TOP_K)


def _fuse(*docsets: list[Document]) -> list[Document]:
    """Union candidate sets, dropping exact duplicates (same source, loc, text)."""
    seen, out = set(), []
    for d in (d for ds in docsets for d in ds):
        key = (d.metadata.get("source"), d.metadata.get("loc"), d.page_content)
        if key not in seen:
            seen.add(key)
            out.append(d)
    return out


def retrieve_docs(query: str, vectorstore: Chroma, embed_key: str) -> list[Document]:
    """Hybrid retrieval: vector ∪ BM25 keyword candidates, reranked down to TOP_K.

    Vector search catches paraphrase/semantic matches; BM25 catches exact terms
    (part numbers, codes, rare words) that embeddings smear. The cross-encoder
    reranker then scores the union against the query and keeps the best TOP_K.
    """
    # ponytail: rebuilds BM25 per query; fine for the ≤2-file session corpus, cache if it grows
    stored = vectorstore.get()
    corpus = [
        Document(page_content=t, metadata=m)
        for t, m in zip(stored["documents"], stored["metadatas"])
    ]
    if not corpus:
        return []
    bm25 = BM25Retriever.from_documents(corpus)
    bm25.k = FETCH_K
    candidates = _fuse(vectorstore.similarity_search(query, k=FETCH_K), bm25.invoke(query))
    try:
        return list(reranker(RERANK_API_KEY or embed_key).compress_documents(candidates, query))
    except Exception:
        log.exception("rerank failed — returning fused candidates")
        return candidates[:TOP_K]


@st.cache_resource
def tavily_client(api_key: str):
    from tavily import TavilyClient

    return TavilyClient(api_key=api_key)


def web_documents(query: str, api_key: str) -> list[Document]:
    """Top web results as Documents, labelled by domain for citations."""
    try:
        res = tavily_client(api_key).search(query, max_results=WEB_RESULTS)
    except Exception:
        log.exception("web search failed")
        return []
    docs = []
    for r in res.get("results", []):
        domain = urlparse(r.get("url", "")).netloc.replace("www.", "") or "web"
        docs.append(
            Document(page_content=r.get("content", ""), metadata={"source": "web", "loc": domain})
        )
    return docs


def summary_stream(paths: list[str], llm):
    """Yield summary tokens for the first SUMMARY_WORDS words of each file."""
    excerpts = []
    for path in paths:
        words = []
        for doc in load_units(path):
            words += doc.page_content.split()
            if len(words) >= SUMMARY_WORDS:
                break
        excerpts.append(f"## {os.path.basename(path)}\n{' '.join(words[:SUMMARY_WORDS])}")
    log.info("LLM request | summarise: %s", [os.path.basename(p) for p in paths])
    chain = SUMMARY_PROMPT | llm | StrOutputParser()
    yield from chain.stream({"excerpts": "\n\n".join(excerpts)})


def answer_stream(
    question: str, history: list[dict], llm, vectorstore: Chroma, tavily_key: str, embed_key: str
):
    """Yield answer tokens from the LCEL chain: retrieve (docs + web) -> prompt -> llm."""
    retrievers = {}
    if vectorstore.get(limit=1)["ids"]:
        retrievers["docs"] = RunnableLambda(lambda q: retrieve_docs(q, vectorstore, embed_key))
    if tavily_key:
        retrievers["web"] = RunnableLambda(lambda q: web_documents(q, tavily_key))
    if not retrievers:
        yield "No documents indexed yet — upload a file, or turn on web search in the sidebar."
        return
    merge = RunnableLambda(lambda hits: format_docs([d for v in hits.values() for d in v]))
    chain = (
        RunnablePassthrough.assign(
            context=itemgetter("question") | RunnableParallel(retrievers) | merge
        )
        | ANSWER_PROMPT
        | llm
        | StrOutputParser()
    )
    log.info("LLM request | Q: %s", question)
    yield from chain.stream({"question": question, "history": history})


def agent_stream(
    question: str, history: list[dict], llm, vectorstore: Chroma, tavily_key: str, embed_key: str
):
    """Agent chat: the LLM picks between document search, web search and the calculator."""
    from langchain.agents import create_agent
    from langchain_core.messages import AIMessageChunk
    from langchain_core.tools import tool

    @tool
    def search_documents(query: str) -> str:
        """Search the user's uploaded documents (PDFs and spreadsheets)."""
        hits = retrieve_docs(query, vectorstore, embed_key)
        return format_docs(hits) if hits else "No documents are indexed yet."

    @tool
    def calculate(expression: str) -> str:
        """Do arithmetic exactly — rates, totals, percentages, GST. E.g. '500 * 12 * 0.18'."""
        try:
            return str(safe_calc(expression))
        except Exception:
            return f"Couldn't evaluate '{expression}' — use plain arithmetic like 500*12*0.18."

    tools = [search_documents, calculate]
    if tavily_key:

        @tool
        def search_web(query: str) -> str:
            """Search the web for general knowledge, logistics concepts or current info."""
            docs = web_documents(query, tavily_key)
            return format_docs(docs) if docs else "No web results found."

        tools.append(search_web)

    agent = create_agent(llm, tools, system_prompt=AGENT_PROMPT)
    log.info("LLM request (agent) | Q: %s", question)

    async def tokens():
        async for chunk, _ in agent.astream(
            {"messages": [*history, {"role": "user", "content": question}]},
            stream_mode="messages",
        ):
            if isinstance(chunk, AIMessageChunk) and isinstance(chunk.content, str) and chunk.content:
                yield chunk.content

    # st.write_stream is sync; drive the async agent one token at a time on a private loop
    gen, loop = tokens(), asyncio.new_event_loop()
    try:
        while True:
            try:
                yield loop.run_until_complete(anext(gen))
            except StopAsyncIteration:
                break
    finally:
        loop.close()


def chat_stream(
    question: str, history: list[dict], llm, vectorstore: Chroma, tavily_key: str, embed_key: str
):
    """Agent (calculator + doc/web search) when the model supports tools; plain chain otherwise."""
    agent = agent_stream(question, history, llm, vectorstore, tavily_key, embed_key)
    try:
        first = next(agent)  # models without tool support error here, before any output
    except StopIteration:
        return
    except AuthenticationError:
        raise  # bad API key — the plain chain would fail the same way, so surface it
    except Exception:
        log.exception("agent unavailable — falling back to the plain RAG chain")
        yield from answer_stream(question, history, llm, vectorstore, tavily_key, embed_key)
        return
    yield first
    yield from agent


def _detect_header(raw: pd.DataFrame) -> int:
    """Real tables often sit below a title/metadata block; find the densest text row.

    Scores the first 20 rows on fill ratio, text share, and how filled the next
    rows are — the header of the actual table wins over 'Client: …' metadata lines.
    """
    best, best_score = 0, -1.0
    width = max(raw.shape[1], 1)
    for i in range(min(len(raw), 20)):
        row = raw.iloc[i]
        filled = int(row.notna().sum())
        if filled < 2:
            continue
        texty = sum(isinstance(v, str) for v in row) / filled
        below = raw.iloc[i + 1 : i + 4]
        below_filled = float(below.notna().to_numpy().mean()) if len(below) else 0.0
        score = filled / width + texty + below_filled
        if score > best_score:
            best, best_score = i, score
    return best


def _coerce_types(col: pd.Series) -> pd.Series:
    """Columns are object dtype after re-heading; restore numbers and dates."""
    if not col.notna().any():
        return col
    threshold = col.notna().sum() * 0.8
    nums = pd.to_numeric(col, errors="coerce")
    if nums.notna().sum() >= threshold:
        return nums
    from datetime import date, datetime

    if col.map(lambda v: isinstance(v, (datetime, date, pd.Timestamp))).sum() >= threshold:
        return pd.to_datetime(col, errors="coerce")
    return col


@st.cache_data(max_entries=8)
def excel_frames(data: bytes) -> dict[str, pd.DataFrame]:
    """All sheets of an uploaded workbook as DataFrames, headers auto-detected."""
    out = {}
    for name, raw in pd.read_excel(io.BytesIO(data), sheet_name=None, header=None).items():
        raw = raw.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
        if raw.empty:
            out[name] = raw
            continue
        h = _detect_header(raw)
        cols, seen = [], {}
        for j, v in enumerate(raw.iloc[h]):
            c = str(v).strip() if pd.notna(v) else f"column {j + 1}"
            seen[c] = seen.get(c, 0) + 1
            cols.append(c if seen[c] == 1 else f"{c} ({seen[c]})")
        df = raw.iloc[h + 1 :].reset_index(drop=True).dropna(how="all")
        df.columns = cols
        out[name] = df.apply(_coerce_types)
    return out


def sheet_chart(df: pd.DataFrame):
    """Auto-chart one sheet: time series -> line, categories -> bar, else index line."""
    df = df.dropna(axis=1, how="all")
    num_cols = list(df.select_dtypes("number").columns)[:CHART_SERIES]
    if not num_cols:
        st.caption("No numeric columns to chart in this sheet.")
        return
    date_cols = [c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])]
    if date_cols:
        st.line_chart(df, x=date_cols[0], y=num_cols)
        return
    cat_cols = [
        c for c in df.columns if c not in num_cols and df[c].dtype == object and df[c].nunique() <= 30
    ]
    if cat_cols:
        top = (
            df.groupby(cat_cols[0])[num_cols]
            .sum()
            .sort_values(num_cols[0], ascending=False)
            .head(12)
        )
        st.bar_chart(top, y=num_cols)
        return
    st.line_chart(df[num_cols])


@st.cache_data(max_entries=32, show_spinner="Designing chart…")
def chart_spec(question: str, columns: str, model_id: str, _llm) -> dict:
    """User's chart request -> validated-later JSON spec (cached per question)."""
    chain = CHART_SPEC_PROMPT | _llm | JsonOutputParser()
    return chain.invoke({"question": question, "columns": columns})


def render_custom_chart(df: pd.DataFrame, spec: dict):
    """Render the chart the spec describes; falls back with a hint when it can't."""
    num = list(df.select_dtypes("number").columns)
    kind = str(spec.get("chart", "bar")).lower()
    x = spec.get("x") if spec.get("x") in df.columns else None
    ys = [c for c in (spec.get("y") or []) if c in num] or num[:1]
    agg = str(spec.get("agg", "sum")).lower()
    top_n = min(int(spec.get("top_n") or 12), 30)

    if kind in ("line", "area", "scatter"):
        if not ys:
            st.warning("No numeric column matched — try naming one from the data preview.")
            return
        {"line": st.line_chart, "area": st.area_chart, "scatter": st.scatter_chart}[kind](
            df, x=x, y=ys
        )
        return
    # bar / pie aggregate by a category
    if not x:
        st.warning("Tell me what to group by — e.g. 'bar of Volume by City'.")
        return
    if agg == "count":
        data = df.groupby(x).size().rename("count").to_frame()
        ys = ["count"]
    elif ys:
        data = df.groupby(x)[ys].agg("mean" if agg == "mean" else "sum")
    else:
        st.warning("No numeric column matched — try naming one from the data preview.")
        return
    data = data.sort_values(ys[0], ascending=False).head(8 if kind == "pie" else top_n)
    if kind == "pie":
        import altair as alt

        st.altair_chart(
            alt.Chart(data.reset_index())
            .mark_arc(innerRadius=45)
            .encode(
                theta=alt.Theta(ys[0], type="quantitative"),
                color=alt.Color(x, type="nominal"),
                tooltip=[x, ys[0]],
            )
        )
        return
    st.bar_chart(data, y=ys)


def render_analysis(files, llm, model_id: str):
    """Charts + previews for uploaded Excel files (analysis toggle)."""
    st.header(":material/query_stats: Data analysis")
    excel = [f for f in files if f.name.lower().endswith((".xlsx", ".xlsm"))]
    if not excel:
        st.caption("Charts need Excel files — PDFs are text-only.")
        return
    for f in excel:
        st.subheader(f.name)
        sheets = excel_frames(f.getvalue())
        tabs = st.tabs(list(sheets))
        for tab, (name, df) in zip(tabs, sheets.items()):
            with tab:
                st.caption(f"{len(df):,} rows · {len(df.columns)} columns")
                request = st.text_input(
                    "Describe a chart",
                    placeholder='e.g. "pie of shipments by status" or "line of Volume over Date"',
                    key=f"chart-{f.name}-{name}",
                    help="XEON AI picks the chart type and columns; leave blank for the automatic chart.",
                )
                if request:
                    cols = ", ".join(f"{c} ({df[c].dtype})" for c in df.columns)
                    try:
                        render_custom_chart(df, chart_spec(request, cols, model_id, llm))
                    except Exception:
                        log.exception("custom chart failed")
                        st.warning("Couldn't build that chart — try naming columns from the data preview.")
                else:
                    sheet_chart(df)
                with st.expander("Data preview", icon=":material/table_rows:"):
                    st.dataframe(df.head(100))


FEEDBACK_LOG = "feedback.jsonl"


def log_feedback(idx: int):
    """Append a thumbs rating for the answer at message `idx` to FEEDBACK_LOG.

    This only *records* signal for later review — it does not change live answers.
    """
    rating = st.session_state.get(f"fb-{idx}")
    if rating is None:  # selection cleared
        return
    msgs = st.session_state.get("messages", [])
    record = {
        "time": datetime.now().isoformat(timespec="seconds"),
        "rating": "up" if rating == 1 else "down",
        "model": st.session_state.get("fb_model", ""),
        "question": msgs[idx - 1]["content"] if 0 < idx < len(msgs) else "",
        "answer": msgs[idx]["content"] if idx < len(msgs) else "",
    }
    try:
        with open(FEEDBACK_LOG, "a", encoding="utf-8") as f:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
        log.info("feedback %s logged", record["rating"])
    except Exception:
        log.exception("feedback log failed")


def main():
    icon = "assets/favicon.png" if os.path.exists("assets/favicon.png") else ":material/description:"
    st.set_page_config(
        page_title="DocMagic",
        page_icon=icon,
        layout="wide",
        menu_items={
            "Get help": None,
            "Report a bug": None,
            "About": "**DocMagic** — chat with your documents, powered by XEON AI.",
        },
    )
    st.session_state.setdefault("messages", [])
    st.session_state.setdefault("summary", "")

    # hide the hover anchor-link icons Streamlit adds to every heading
    st.markdown(
        "<style>[data-testid='stHeaderActionElements']{display:none}</style>",
        unsafe_allow_html=True,
    )

    with st.sidebar:
        st.subheader(":material/key: Access")
        pasted = st.text_input(
            "Chat API key",
            type="password",
            placeholder="Enter your API key",
            help="Bring your own free NVIDIA key for chat (build.nvidia.com). "
            "Document search runs on the app's key — nothing needed from you.",
        )
        model = st.selectbox(
            "Chat model",
            list(dict.fromkeys([LLM_MODEL, *CURATED_MODELS])) + [CUSTOM_MODEL],
            help="Models served via NVIDIA NIM.",
        )
        if model == CUSTOM_MODEL:
            model = st.text_input("Model id", placeholder="provider/model-name").strip() or LLM_MODEL
        with st.expander("Advanced", icon=":material/tune:"):
            base_url = st.text_input("LLM base URL", value=LLM_BASE_URL)

        st.subheader(":material/folder: Documents")
        files = st.file_uploader(
            f"Up to {MAX_FILES} files",
            type=["pdf", "xlsx", "xlsm"],
            accept_multiple_files=True,
        )
        analyse = st.button(
            "Submit & analyse", type="primary", icon=":material/auto_awesome:", width="stretch"
        )
        analysis_on = st.toggle(
            "Enable analysis", help="Charts and data previews for uploaded Excel files"
        )
        has_web = bool(secret("TAVILY_API_KEY"))
        web_on = st.toggle(
            "Search the web",
            value=has_web,
            disabled=not has_web,
            help="Answer from live web results (Tavily) alongside your documents. "
            "Needs TAVILY_API_KEY in secrets or .env.",
        )

    st.title("DocMagic")
    st.caption(
        "Chat with your documents, powered by XEON AI — answers cited to the exact file, page or sheet."
    )

    llm_key = pasted or secret("LLM_API_KEY")
    # embeddings run on the app's own key, so document search and the knowledge
    # base stay free; only chatting and summarising need a user's chat key.
    embed_key = secret("EMBED_API_KEY") or pasted or llm_key
    if not embed_key:
        st.info(
            "Enter an NVIDIA API key in the sidebar to begin — a free one at build.nvidia.com.",
            icon=":material/key:",
        )
        st.stop()
    llm, embeddings = resources(llm_key or "no-chat-key", base_url, model, embed_key)
    vectorstore = docs_store(embeddings)
    tavily_key = secret("TAVILY_API_KEY") if web_on else ""

    # side-by-side layout when a PDF is uploaded: chat on the left, live preview on the right
    pdfs = [f for f in (files or []) if f.name.lower().endswith(".pdf")]
    main_col, preview_col = st.columns([3, 2], gap="large") if pdfs else (st.container(), None)

    with main_col:
        if analyse:
            if not files:
                st.warning("Upload at least one file first.")
            elif len(files) > MAX_FILES:
                st.warning(f"Please upload at most {MAX_FILES} files at a time.")
            else:
                try:
                    # uploads are in-memory; loaders want paths, so spill to a temp dir
                    with tempfile.TemporaryDirectory() as td:
                        paths = []
                        for f in files:
                            p = os.path.join(td, f.name)
                            with open(p, "wb") as out:
                                out.write(f.getbuffer())
                            paths.append(p)
                        with st.spinner("Reading your documents…"):
                            # each Submit & analyse starts fresh so chat only covers current files
                            vectorstore.reset_collection()
                            n = ingest(paths, vectorstore)
                        st.session_state.messages = []  # old chat referred to old docs
                        if n == 0:
                            st.session_state.summary = ""
                            st.warning(
                                "I couldn't find readable text in those files — a scanned or "
                                "image-only PDF has no text layer to read.",
                                icon=":material/document_scanner:",
                            )
                        elif not llm_key:
                            st.session_state.summary = ""
                            st.success(
                                "Documents indexed. Add your API key in the sidebar to chat "
                                "and get a summary.",
                                icon=":material/check_circle:",
                            )
                        else:
                            st.toast("Documents ready", icon=":material/check_circle:")
                            with st.expander(
                                "Document summary", icon=":material/description:", expanded=True
                            ):
                                st.session_state.summary = st.write_stream(summary_stream(paths, llm))
                except Exception:
                    log.exception("analyse failed")
                    st.error(
                        "Couldn't read those documents — check the files and try again.",
                        icon=":material/error:",
                    )
        elif st.session_state.summary:
            with st.expander("Document summary", icon=":material/description:"):
                st.markdown(st.session_state.summary)

        if analysis_on and files:
            render_analysis(files, llm, model)

        st.session_state["fb_model"] = model  # recorded alongside any feedback
        for i, m in enumerate(st.session_state.messages):
            with st.chat_message(m["role"]):
                st.markdown(m["content"])
                if m["role"] == "assistant":
                    st.feedback("thumbs", key=f"fb-{i}", on_change=log_feedback, args=(i,))

        question = None
        if not st.session_state.messages:
            picked = st.pills("Try asking", list(SUGGESTIONS), label_visibility="collapsed")
            if picked:
                question = SUGGESTIONS[picked]

    if preview_col is not None:
        with preview_col:
            st.subheader(":material/picture_as_pdf: Document preview")
            choice = pdfs[0]
            if len(pdfs) > 1:
                name = st.selectbox("File", [f.name for f in pdfs], label_visibility="collapsed")
                choice = next(f for f in pdfs if f.name == name)
            st.pdf(choice.getvalue(), height=640)

    prompt = st.chat_input("Ask about your documents", submit_mode="disable")
    question = prompt or question

    if question and not llm_key:
        with main_col:
            st.chat_message("user").markdown(question)
            st.info(
                "Add your NVIDIA API key in the sidebar to chat about your documents.",
                icon=":material/key:",
            )
    elif question:
        with main_col:
            st.chat_message("user").markdown(question)
            try:
                with st.chat_message("assistant"):
                    text = st.write_stream(
                        chat_stream(
                            question, st.session_state.messages, llm, vectorstore, tavily_key, embed_key
                        )
                    )
            except AuthenticationError:  # bad/rejected chat key — actionable message
                log.exception("chat key rejected (401)")
                st.error(
                    "Your chat API key was rejected. Check it's a valid key for the selected "
                    "model (NVIDIA keys start with `nvapi-`) and try again.",
                    icon=":material/key_off:",
                )
            except Exception:  # rate limits & co: tell the user, keep history clean
                log.exception("answer failed")
                st.error(
                    "Sorry, I hit a snag answering that — please try again.", icon=":material/error:"
                )
            else:
                st.session_state.messages.extend(
                    [{"role": "user", "content": question}, {"role": "assistant", "content": text}]
                )
                st.rerun()  # re-render so the new answer shows its feedback buttons

    # sidebar status LAST, so the document count reflects this run's ingest
    with st.sidebar:
        docs_ready = "ready" if vectorstore._collection.count() else "none yet"
        st.caption(f"Your documents: {docs_ready} · Web search: {'on' if web_on else 'off'}")


if __name__ == "__main__":
    main()
