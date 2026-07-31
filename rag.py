"""DocMagic core — config, models, retrieval, storage and tool primitives.

Pure library: no web framework, no UI. api.py (FastAPI) and the LangGraph
agent/ingestion graphs all import from here.
"""

import ast
import io
import logging
import operator
import os
from functools import lru_cache
from hashlib import md5
from urllib.parse import urlparse

import pandas as pd
from dotenv import load_dotenv
from langchain_community.document_loaders import PyPDFLoader
from langchain_community.retrievers import BM25Retriever
from langchain_core.documents import Document
from langchain_nvidia_ai_endpoints import NVIDIAEmbeddings, NVIDIARerank
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy import create_engine, text

load_dotenv(override=True)

# root stays at WARNING so httpx/uvicorn chatter is hidden; only our logs show
logging.basicConfig(format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("docmagic")
log.setLevel(logging.INFO)

LLM_BASE_URL = os.getenv("LLM_BASE_URL", "https://integrate.api.nvidia.com/v1")
LLM_API_KEY = os.getenv("LLM_API_KEY", "")
LLM_MODEL = os.getenv("LLM_MODEL", "auto")  # the gateway's router picks the model
EMBED_API_KEY = os.getenv("EMBED_API_KEY") or LLM_API_KEY
# embeddings keep their own endpoint so switching LLM provider can't break them
EMBED_BASE_URL = os.getenv("EMBED_BASE_URL", "https://integrate.api.nvidia.com/v1")
EMBED_MODEL = os.getenv("EMBED_MODEL", "nvidia/nemotron-3-embed-1b")
RERANK_MODEL = os.getenv("RERANK_MODEL", "nvidia/llama-nemotron-rerank-vl-1b-v2")  # reranks hybrid hits
RERANK_API_KEY = os.getenv("RERANK_API_KEY", "")  # own key; falls back to the embed key
POSTGRES_URL = os.getenv("POSTGRES_URL", "")  # required: pgvector chunks + extracted sheets

CHUNK_CHARS = 1500  # splitter counts characters; ~250 words
CHUNK_OVERLAP = 300  # ~50 words
TOP_K = 5
FETCH_K = 20  # hybrid candidates pulled per retriever before reranking down to TOP_K
CORPUS_K = 500  # ceiling for the per-query BM25 corpus pull (a session holds far fewer chunks)
WEB_RESULTS = 4  # web results pulled per question when web search is on

SYSTEM_PROMPT = (
    "You are XEON AI, a friendly expert on logistics, warehousing, ERP and CRM "
    "documents (SOPs, rate cards, contracts, manuals). Talk like a helpful "
    "colleague: warm, plain language, short sentences — never stiff or robotic. "
    "For questions about the documents, answer ONLY from the provided context and "
    "cite the source like [file.pdf p.3] or [file.xlsx Rates] (sheet name for "
    "spreadsheets). The context may include live web results labeled [web ...]; "
    "prefer the user's documents for specifics and use web results for general or "
    "current information, citing them by domain. If the context does not "
    "have the answer, say so in one friendly sentence — do not lecture about what "
    "the context contains. "
    "If the user is just greeting or reacting ('hi', 'nice', 'thanks'), reply "
    "warmly in a sentence or two — no citations, no mention of context or documents."
)

AGENT_PROMPT = SYSTEM_PROMPT + (
    " You have tools. Use search_documents for anything about the user's uploaded files. "
    "Use search_web for general knowledge, logistics concepts or current information. "
    "Use calculate for ALL arithmetic — rates, totals, percentages, GST — never do math "
    "in your head; fetch the numbers first, then calculate. Show the working briefly and "
    "cite document excerpts by file and web results by domain."
)

# splits on paragraphs first, then lines, then words — chunks end at natural
# boundaries instead of mid-sentence
splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_CHARS, chunk_overlap=CHUNK_OVERLAP)


@lru_cache(maxsize=4)
def resources(llm_key: str, llm_base_url: str, llm_model: str, embed_key: str):
    """LLM + embeddings, cached per (key, model) so a model switch takes effect."""
    kwargs = {}
    if "nemotron" in llm_model:
        kwargs["extra_body"] = {"chat_template_kwargs": {"enable_thinking": False}}
    if "kimi" not in llm_model:
        kwargs["temperature"] = 0.2  # kimi-k3 fixes temperature=1.0 and rejects overrides
    llm = ChatOpenAI(base_url=llm_base_url, api_key=llm_key, model=llm_model, **kwargs)
    # embed_documents() sends input_type=passage, embed_query() sends query
    embeddings = NVIDIAEmbeddings(
        base_url=EMBED_BASE_URL, model=EMBED_MODEL, api_key=embed_key, truncate="END"
    )
    return llm, embeddings


def default_resources():
    """Env-configured resources, for scripts and tests."""
    return resources(LLM_API_KEY, LLM_BASE_URL, LLM_MODEL, EMBED_API_KEY)


# in-process calculator (Cloud-safe: no subprocess, no MCP server to spawn)
_CALC_OPS = {
    ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
    ast.Div: operator.truediv, ast.FloorDiv: operator.floordiv, ast.Mod: operator.mod,
    ast.Pow: operator.pow, ast.USub: operator.neg, ast.UAdd: operator.pos,
}


def safe_calc(expression: str) -> float:
    """Evaluate arithmetic safely — numbers and + - * / // % ** only, never runs code."""

    def _eval(node):
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.BinOp) and type(node.op) in _CALC_OPS:
            return _CALC_OPS[type(node.op)](_eval(node.left), _eval(node.right))
        if isinstance(node, ast.UnaryOp) and type(node.op) in _CALC_OPS:
            return _CALC_OPS[type(node.op)](_eval(node.operand))
        raise ValueError("unsupported expression")

    return _eval(ast.parse(expression, mode="eval").body)


@lru_cache(maxsize=1)
def engine():
    """SQLAlchemy engine, with the app's own schema created on first use.

    `analytics` holds the sheets extracted from uploads (what the data specialist
    queries and promptable Vega-Lite charts) plus the feedback log. pgvector manages
    its own tables separately.
    """
    if not POSTGRES_URL:
        raise RuntimeError("POSTGRES_URL is required - set it in .env (see .env.example)")
    # pool_pre_ping: hosted Postgres (Supabase/Neon) drops idle connections
    eng = create_engine(POSTGRES_URL, pool_pre_ping=True)
    with eng.begin() as c:
        c.execute(text("CREATE SCHEMA IF NOT EXISTS analytics"))
        c.execute(
            text(
                "CREATE TABLE IF NOT EXISTS analytics.uploads ("
                "sid text, source text, sheet text, table_name text, rows integer, "
                "created timestamptz DEFAULT now(), PRIMARY KEY (sid, source, sheet))"
            )
        )
        c.execute(
            text(
                "CREATE TABLE IF NOT EXISTS analytics.feedback ("
                "id bigserial PRIMARY KEY, created timestamptz DEFAULT now(), sid text, "
                "rating text, model text, question text, answer text)"
            )
        )
    return eng


def docs_store(embeddings, sid: str):
    """Per-session pgvector collection: visitors never see each other's files."""
    if not POSTGRES_URL:
        raise RuntimeError("POSTGRES_URL is required - set it in .env (see .env.example)")
    from langchain_postgres import PGVector

    # ponytail: orphaned session collections persist; sweep stale docs-* if the demo DB fills
    return PGVector(
        embeddings=embeddings,
        collection_name=f"docs-{sid}",
        connection=POSTGRES_URL,
        use_jsonb=True,
    )


def reset_store(vectorstore) -> None:
    """Empty the session's vector collection before a fresh ingest."""
    vectorstore.delete_collection()
    vectorstore.create_collection()


def _table_name(sid: str, source: str, sheet: str) -> str:
    """Deterministic, always-valid table name — sheet names can be anything."""
    return "t_" + md5(f"{sid}|{source}|{sheet}".encode()).hexdigest()[:16]


def save_frames(sid: str, source: str, frames: dict) -> None:
    """Extracted sheets → one Postgres table each, registered in analytics.uploads."""
    eng = engine()
    for sheet, df in frames.items():
        if df.empty:
            continue
        table = _table_name(sid, source, sheet)
        df.to_sql(table, eng, schema="analytics", if_exists="replace", index=False)
        with eng.begin() as c:
            c.execute(
                text(
                    "INSERT INTO analytics.uploads (sid, source, sheet, table_name, rows) "
                    "VALUES (:sid, :source, :sheet, :table, :rows) "
                    "ON CONFLICT (sid, source, sheet) DO UPDATE SET "
                    "table_name = EXCLUDED.table_name, rows = EXCLUDED.rows"
                ),
                {"sid": sid, "source": source, "sheet": sheet, "table": table, "rows": len(df)},
            )
        log.info("extracted %s · %s → analytics.%s (%d rows)", source, sheet, table, len(df))


def reset_uploads(sid: str) -> None:
    """Drop this session's extracted tables before a fresh ingest."""
    eng = engine()
    with eng.begin() as c:
        tables = c.execute(
            text("SELECT table_name FROM analytics.uploads WHERE sid = :sid"), {"sid": sid}
        ).scalars().all()
        for t in tables:  # names are our own md5 hex, never user input — safe to interpolate
            c.execute(text(f'DROP TABLE IF EXISTS analytics."{t}"'))
        c.execute(text("DELETE FROM analytics.uploads WHERE sid = :sid"), {"sid": sid})


def sheet_names(sid: str) -> list[str]:
    """Sheets this session has extracted (empty when nothing tabular was uploaded)."""
    with engine().begin() as c:
        return list(
            c.execute(
                text("SELECT sheet FROM analytics.uploads WHERE sid = :sid"), {"sid": sid}
            ).scalars()
        )


def load_frame(sid: str, sheet: str) -> pd.DataFrame | None:
    """One extracted sheet back as a DataFrame (case-insensitive lookup)."""
    with engine().begin() as c:
        table = c.execute(
            text(
                "SELECT table_name FROM analytics.uploads "
                "WHERE sid = :sid AND lower(sheet) = lower(:sheet)"
            ),
            {"sid": sid, "sheet": sheet},
        ).scalar()
    return None if table is None else pd.read_sql_table(table, engine(), schema="analytics")


def load_units(path: str) -> list[Document]:
    """One Document per PDF page or Excel sheet, labeled for citations."""
    from openpyxl import load_workbook

    name = os.path.basename(path)
    if path.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(path, read_only=True, data_only=True)
        docs = []
        for ws in wb.worksheets:
            rows = (
                " | ".join(str(c) for c in row if c is not None)
                for row in ws.iter_rows(values_only=True)
            )
            text = "\n".join(r for r in rows if r)
            docs.append(Document(page_content=text, metadata={"source": name, "loc": ws.title}))
        wb.close()
        return docs
    return [
        Document(
            page_content=p.page_content,
            metadata={"source": name, "loc": f"p.{p.metadata.get('page', 0) + 1}"},
        )
        for p in PyPDFLoader(path).load()
    ]


def ingest(paths: list[str], vectorstore) -> int:
    """Load, split and index the documents; returns the chunk count."""
    chunks = splitter.split_documents(doc for path in paths for doc in load_units(path))
    if chunks:
        vectorstore.add_documents(chunks)  # NVIDIAEmbeddings batches 50/call itself
    return len(chunks)


def format_docs(docs: list[Document]) -> str:
    log.info("context <- %s", [f"{d.metadata['source']} {d.metadata['loc']}" for d in docs])
    return "\n\n".join(
        f"[{d.metadata['source']} {d.metadata['loc']}]\n{d.page_content}" for d in docs
    )


@lru_cache(maxsize=2)
def reranker(api_key: str):
    """NVIDIA NIM cross-encoder reranker; runs on the app's free embed key/endpoint."""
    return NVIDIARerank(model=RERANK_MODEL, base_url=EMBED_BASE_URL, api_key=api_key, top_n=TOP_K)


def _fuse(*docsets: list[Document]) -> list[Document]:
    """Union candidate sets, dropping exact duplicates (same source, loc, text)."""
    seen, out = set(), []
    for d in (d for ds in docsets for d in ds):
        key = (d.metadata.get("source"), d.metadata.get("loc"), d.page_content)
        if key not in seen:
            seen.add(key)
            out.append(d)
    return out


def retrieve_docs(query: str, vectorstore, embed_key: str, progress=None) -> list[Document]:
    """Hybrid retrieval: vector ∪ BM25 keyword candidates, reranked down to TOP_K.

    Vector search catches paraphrase/semantic matches; BM25 catches exact terms
    (part numbers, codes, rare words) that embeddings smear. The cross-encoder
    reranker then scores the union against the query and keeps the best TOP_K.

    `progress(stage)` is called as each step starts, so the UI can narrate the real
    pipeline instead of one generic label. It may be called from a worker thread.
    """
    step = progress or (lambda _stage: None)
    # ponytail: one big similarity_search doubles as the vector hits AND the BM25 corpus —
    # PGVector has no "fetch the whole collection" call, and CORPUS_K covers a session's chunks
    step("embedding")
    corpus = vectorstore.similarity_search(query, k=CORPUS_K)
    if not corpus:
        return []
    step("keyword")
    bm25 = BM25Retriever.from_documents(corpus)
    bm25.k = FETCH_K
    candidates = _fuse(corpus[:FETCH_K], bm25.invoke(query))
    step("rerank")
    try:
        return list(reranker(RERANK_API_KEY or embed_key).compress_documents(candidates, query))
    except Exception:
        log.exception("rerank failed — returning fused candidates")
        return candidates[:TOP_K]


@lru_cache(maxsize=2)
def tavily_client(api_key: str):
    from tavily import TavilyClient

    return TavilyClient(api_key=api_key)


def web_documents(query: str, api_key: str) -> list[Document]:
    """Top web results as Documents, labelled by domain for citations."""
    try:
        res = tavily_client(api_key).search(query, max_results=WEB_RESULTS)
    except Exception:
        log.exception("web search failed")
        return []
    docs = []
    for r in res.get("results", []):
        domain = urlparse(r.get("url", "")).netloc.replace("www.", "") or "web"
        docs.append(
            Document(page_content=r.get("content", ""), metadata={"source": "web", "loc": domain})
        )
    return docs


@lru_cache(maxsize=1)
def langfuse_handler():
    """Langfuse tracing callback, or None if keys are unset or the SDK can't init.

    v4's CallbackHandler reads creds from the global client (env vars
    LANGFUSE_PUBLIC_KEY / _SECRET_KEY / _HOST or _BASE_URL) — no args needed.
    """
    if not (os.getenv("LANGFUSE_PUBLIC_KEY") and os.getenv("LANGFUSE_SECRET_KEY")):
        return None
    try:
        from langfuse.langchain import CallbackHandler

        return CallbackHandler()
    except Exception:
        log.exception("Langfuse tracing disabled — handler init failed")
        return None


def trace_config() -> dict:
    """LangChain run config that attaches Langfuse tracing when it's configured."""
    handler = langfuse_handler()
    return {"callbacks": [handler]} if handler else {}


def _detect_header(raw: pd.DataFrame) -> int:
    """Real tables often sit below a title/metadata block; find the densest text row.

    Scores the first 20 rows on fill ratio, text share, and how filled the next
    rows are — the header of the actual table wins over 'Client: …' metadata lines.
    """
    best, best_score = 0, -1.0
    width = max(raw.shape[1], 1)
    for i in range(min(len(raw), 20)):
        row = raw.iloc[i]
        filled = int(row.notna().sum())
        if filled < 2:
            continue
        texty = sum(isinstance(v, str) for v in row) / filled
        below = raw.iloc[i + 1 : i + 4]
        below_filled = float(below.notna().to_numpy().mean()) if len(below) else 0.0
        score = filled / width + texty + below_filled
        if score > best_score:
            best, best_score = i, score
    return best


def _coerce_types(col: pd.Series) -> pd.Series:
    """Columns are object dtype after re-heading; restore numbers and dates."""
    if not col.notna().any():
        return col
    threshold = col.notna().sum() * 0.8
    nums = pd.to_numeric(col, errors="coerce")
    if nums.notna().sum() >= threshold:
        return nums
    from datetime import date, datetime

    if col.map(lambda v: isinstance(v, (datetime, date, pd.Timestamp))).sum() >= threshold:
        return pd.to_datetime(col, errors="coerce")
    return col


# ponytail: lru_cache shares the frames (st.cache_data copied); no caller mutates them.
# maxsize stays small — each entry holds a whole workbook's DataFrames in memory.
@lru_cache(maxsize=4)
def excel_frames(data: bytes) -> dict[str, pd.DataFrame]:
    """All sheets of an uploaded workbook as DataFrames, headers auto-detected."""
    out = {}
    for name, raw in pd.read_excel(io.BytesIO(data), sheet_name=None, header=None).items():
        raw = raw.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
        if raw.empty:
            out[name] = raw
            continue
        h = _detect_header(raw)
        cols, seen = [], {}
        for j, v in enumerate(raw.iloc[h]):
            c = str(v).strip() if pd.notna(v) else f"column {j + 1}"
            seen[c] = seen.get(c, 0) + 1
            cols.append(c if seen[c] == 1 else f"{c} ({seen[c]})")
        df = raw.iloc[h + 1 :].reset_index(drop=True).dropna(how="all")
        df.columns = cols
        out[name] = df.apply(_coerce_types)
    return out
